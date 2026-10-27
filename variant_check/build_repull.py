# -*- coding: utf-8 -*-
"""把重抓结果整理成交付表（变体组内 + 未捆绑 两张清单 + 家族表 + 与上次对比）。
用法: python build_repull.py --json repull_all.json --asins asins.txt --prev 旧交付.xlsx --out 新交付.xlsx
"""
import argparse
import json
import re
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def num(s):
    m = re.search(r"([\d.]+)", str(s if s is not None else ""))
    return float(m.group(1)) if m else None


def cnt(s):
    m = re.search(r"([\d,]+)", str(s if s is not None else ""))
    return int(m.group(1).replace(",", "")) if m else None


class UF:
    def __init__(self):
        self.p = {}

    def find(self, x):
        self.p.setdefault(x, x)
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.p[rb] = ra


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True)
    ap.add_argument("--asins", required=True)
    ap.add_argument("--prev")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    today = str(date.today())

    asins = [x.strip() for x in Path(a.asins).read_text(encoding="utf-8").replace("\n", ",").split(",") if x.strip()]
    rows = json.loads(Path(a.json).read_text(encoding="utf-8"))["rows"]

    prev = {}
    if a.prev and Path(a.prev).exists():
        xl = pd.ExcelFile(a.prev)
        for sh in ("捆绑清单", "未捆绑清单"):
            if sh in xl.sheet_names:
                df = pd.read_excel(xl, sh)
                for _, r in df.iterrows():
                    prev[str(r["ASIN"]).strip()] = ("在变体组内" if r["前台是否捆绑变体"] == "是" else "未捆绑",
                                                   r.get("家族编号"), r.get("前台家族ASIN数"))

    rec, mem = [], {}
    for sid in asins:
        d = rows.get(sid)
        if not d:
            rec.append({"ASIN": sid, "前台状态": "未抓到", "前台家族ASIN数": None,
                        "前台落地ASIN": "", "变体维度": "", "变体值示例": "", "同家族我方ASIN数": None,
                        "同家族清单内其他ASIN": "", "家族成员ASIN": "", "前台星级": None,
                        "前台评分人数": None, "在售状态": "", "品牌": "", "标题": "",
                        "说明": "抓取失败", "抓取时间": today})
            continue
        n = max(int(d.get("nStateAsins") or 0), int(d.get("nDomAsins") or 0))
        members = set(x for x in (d.get("stateAsins") or []) if x)
        members |= set(x for x in (d.get("domAsins") or []) if x)
        n = max(n, len(members))
        vals = (d.get("vals") or [])[:10]
        note = []
        if n == 1:
            only = next(iter(members), "")
            note.append("有选择器但仅1个选项(%s)" % (("指向 %s" % only) if only and only != sid else "即自身"))
            if only and only != sid:
                note.append("⚠️残留在他人变体组")
        elif n == 0:
            note.append("前台无变体选择器")
        if d.get("landedAsin") and d["landedAsin"] != sid:
            note.append("前台落地 %s" % d["landedAsin"])
        rec.append({"ASIN": sid, "前台状态": "在变体组内" if n >= 2 else "未捆绑",
                    "前台家族ASIN数": max(n, 1),
                    "前台落地ASIN": d.get("landedAsin", ""),
                    "变体维度": "/".join(d.get("dimN") or []) or "",
                    "变体值示例": " ｜ ".join(vals[:8]),
                    "同家族我方ASIN数": None,
                    "同家族清单内其他ASIN": "", "家族成员ASIN": "/".join(sorted(members)),
                    "前台星级": num(d.get("rating")), "前台评分人数": cnt(d.get("ratingCount")),
                    "在售状态": d.get("inStock", ""),
                    "品牌": (d.get("brand") or "").replace("Visit the ", "").replace(" Store", ""),
                    "标题": d.get("title", ""), "说明": "；".join(note) or "正常",
                    "抓取时间": today})
        if members:
            mem[sid] = members

    det = pd.DataFrame(rec)
    det["前台家族ASIN数"] = pd.to_numeric(det["前台家族ASIN数"], errors="coerce")

    uf = UF()
    for sid, m in mem.items():
        uf.find(sid)
        for x in m:
            uf.union(sid, x)
    comp = {}
    for sid in det["ASIN"]:
        comp.setdefault(uf.find(sid), []).append(sid)
    bundled_asins = set(det.loc[det["前台状态"] == "在变体组内", "ASIN"])
    bundled = [(r, o) for r, o in comp.items() if len(set(o) & bundled_asins) >= 2]
    bundled.sort(key=lambda kv: -len(set(kv[1]) & bundled_asins))
    famno, fammem = {}, {}
    for i, (root, ours) in enumerate(bundled, 1):
        mm = set()
        for x in ours:
            mm |= mem.get(x, set())
        famno[root] = i
        fammem[root] = mm

    bundled_asins = set(det.loc[det["前台状态"] == "在变体组内", "ASIN"])
    det["家族编号"] = det["ASIN"].apply(
        lambda s: famno.get(uf.find(s), ("—(与清单外ASIN同组)" if s in bundled_asins else "")))
    det["同家族我方ASIN数"] = det["ASIN"].apply(
        lambda s: len(set(comp.get(uf.find(s), [s])) & bundled_asins) or None)
    det["同家族清单内其他ASIN"] = det["ASIN"].apply(
        lambda s: "/".join(sorted([x for x in set(comp.get(uf.find(s), [])) if x != s and x in bundled_asins])))
    det["家族成员ASIN"] = det["ASIN"].apply(
        lambda s: "/".join(sorted(fammem.get(uf.find(s), mem.get(s, set())))))
    det["上次判定"] = det["ASIN"].apply(lambda s: prev.get(s, ("", "", ""))[0])
    det["上次家族编号"] = det["ASIN"].apply(lambda s: prev.get(s, ("", "", ""))[1])

    def chg(r):
        p = prev.get(r["ASIN"], ("", "", ""))[0]
        if not p:
            return "新抓取"
        return "无变化" if p == r["前台状态"] else "有变化(%s→%s)" % (p, r["前台状态"])

    det["与上次对比"] = det.apply(chg, axis=1)
    det["首次?凑数"] = ""

    cols = ["ASIN", "前台状态", "家族编号", "前台家族ASIN数", "同家族我方ASIN数",
            "同家族清单内其他ASIN", "家族成员ASIN", "变体维度", "变体值示例", "前台星级",
            "前台评分人数", "在售状态", "品牌", "标题", "说明", "上次判定", "与上次对比", "抓取时间"]
    det = det[[c for c in cols if c in det.columns]]
    det = det.sort_values(["前台状态", "家族编号", "ASIN"], ascending=[True, True, True]).reset_index(drop=True)
    inb = det[det["前台状态"] == "在变体组内"].copy()
    out = det[det["前台状态"] == "未捆绑"].copy()

    fams = []
    for root, ours in sorted(bundled, key=lambda kv: famno[kv[0]]):
        ours = sorted(set(ours) & bundled_asins)
        if len(ours) < 2:
            continue
        sub = det[det["ASIN"].isin(ours)]
        n = sub["前台家族ASIN数"].max()
        n = int(n) if pd.notna(n) else len(ours)
        mm = sorted(fammem[root])
        fams.append({"家族编号": famno[root], "前台家族ASIN数": max(n, len(mm)),
                     "清单内我方ASIN数": len(ours),
                     "超出清单的成员数": max(max(n, len(mm)) - len(ours), 0),
                     "变体维度": next((x for x in sub["变体维度"] if x), ""),
                     "清单内我方ASIN": "/".join(ours),
                     "家族成员ASIN": "/".join(mm)[:1200]})
    famdf = pd.DataFrame(fams)

    srows = [["重抓 ASIN 总数", len(det), "口径：前台变体选择器可选的不同 ASIN 数"],
             ["在变体组内", len(inb), "可变体选项 ≥2"],
             ["未捆绑（单ASIN）", len(out), "无选择器 / 仅1个选项"],
             ["变体家族数", len(famdf), "家族内我方链接 ≥2 条"],
             ["最大家族选项数", int(famdf["前台家族ASIN数"].max()) if len(famdf) else 0, ""],
             ["与上次判定有变化", int(det["与上次对比"].str.startswith("有变化").sum()), ""],
             ["抓取日期", today, "amazon.com 前台实测"]]
    summ = pd.DataFrame(srows, columns=["指标", "数值", "说明"])

    chg = det[det["与上次对比"].str.startswith("有变化")][
        ["ASIN", "前台状态", "上次判定", "与上次对比", "家族编号", "前台家族ASIN数",
         "同家族清单内其他ASIN", "说明", "标题"]].sort_values(
        ["上次判定", "ASIN"]).reset_index(drop=True)

    with pd.ExcelWriter(a.out, engine="openpyxl") as xw:
        summ.to_excel(xw, sheet_name="汇总", index=False)
        inb.to_excel(xw, sheet_name="在变体组内", index=False)
        out.to_excel(xw, sheet_name="未捆绑", index=False)
        famdf.to_excel(xw, sheet_name="变体家族", index=False)
        chg.to_excel(xw, sheet_name="与上次变化", index=False)
        hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5597")
        thin = Side(style="thin", color="BFBFBF")
        bd = Border(left=thin, right=thin, top=thin, bottom=thin)
        lefty = ("ASIN", "标题", "说明", "家族成员ASIN", "同家族清单内其他ASIN", "清单内我方ASIN", "变体值示例")
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font, c.fill, c.border = hf, hfill, bd
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 28
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
                nm = str(col[0].value or "")
                if nm == "标题":
                    w = 42
                elif "成员ASIN" in nm or "其他ASIN" in nm or "我方ASIN" in nm or "变体值" in nm:
                    w = 38
                elif "说明" in nm:
                    w = 36
                elif "ASIN" in nm:
                    w = 15
                else:
                    w = 12
                ws.column_dimensions[get_column_letter(j)].width = w
            for row in ws.iter_rows(min_row=2):
                for i, c in enumerate(row):
                    c.border = bd
                    hdr = str(ws.cell(row=1, column=i + 1).value or "")
                    c.alignment = Alignment(horizontal="left" if any(k in hdr for k in lefty) else "center",
                                            vertical="center", wrap_text=True)
            if ws.title != "汇总":
                ws.auto_filter.ref = ws.dimensions
    print("总数 %d | 在变体组内 %d | 未捆绑 %d | 家族 %d | 有变化 %d" % (
        len(det), len(inb), len(out), len(famdf), int(det["与上次对比"].str.startswith("有变化").sum())))
    print("输出:", a.out)


if __name__ == "__main__":
    main()
